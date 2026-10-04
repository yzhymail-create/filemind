#!/usr/bin/env python3
"""内容解析层:把各格式文件提成"段落列表 + 全文 + 提取状态"."""
from __future__ import annotations

import os
from dataclasses import dataclass, field


@dataclass
class ExtractResult:
    paragraphs: list = field(default_factory=list)
    text: str = ""
    status: str = "ok"      # ok | partial | failed
    reason: str = ""


TEXT_EXTS = {".txt", ".md", ".csv", ".log", ".ini", ".cfg", ".json", ".xml"}


def _read_text_file(path: str) -> ExtractResult:
    raw = open(path, "rb").read()
    if not raw.strip():
        return ExtractResult([], "", "ok", "")
    for enc in ("utf-8", "gbk", "gb2312", "big5"):
        try:
            s = raw.decode(enc)
            paras = [p for p in s.splitlines() if p.strip()]
            if enc == "utf-8":
                return ExtractResult(paras, s, "ok", "")
            return ExtractResult(paras, s, "partial", "非 UTF-8 编码,按 %s 解码" % enc)
        except UnicodeDecodeError:
            continue
    s = raw.decode("latin1")
    paras = [p for p in s.splitlines() if p.strip()]
    return ExtractResult(paras, s, "partial", "未知编码,按 latin1 兜底解码")


def _extract_docx(path: str) -> ExtractResult:
    from docx import Document
    doc = Document(path)
    paras = []
    for p in doc.paragraphs:
        t = p.text.strip()
        if t:
            paras.append(t)
    for tbl in doc.tables:
        for row in tbl.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                paras.append(" | ".join(cells))
    for sec in doc.sections:
        for part in (sec.header, sec.footer):
            for p in part.paragraphs:
                t = p.text.strip()
                if t and t not in paras:
                    paras.append(t)
    text = "\n".join(paras)
    if paras:
        return ExtractResult(paras, text, "ok", "")
    return ExtractResult([], "", "partial", "未提取到正文")


def _extract_xlsx(path: str) -> ExtractResult:
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    paras = []
    for ws in wb.worksheets:
        lines = ["【工作表:%s】" % ws.title]
        for row in ws.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c not in (None, "")]
            if cells:
                lines.append(" | ".join(cells))
        if len(lines) > 1:
            paras.extend(lines)
    wb.close()
    text = "\n".join(paras)
    if text.strip():
        return ExtractResult(paras, text, "ok", "")
    return ExtractResult([], "", "partial", "工作表为空或仅含公式")


def _extract_pptx(path: str) -> ExtractResult:
    from pptx import Presentation
    prs = Presentation(path)
    paras = []
    for i, slide in enumerate(prs.slides, 1):
        lines = ["【第%d页】" % i]
        for shape in slide.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    t = p.text.strip()
                    if t:
                        lines.append(t)
            if shape.has_table:
                for row in shape.table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        lines.append(" | ".join(cells))
        if len(lines) > 1:
            paras.extend(lines)
    text = "\n".join(paras)
    if text.strip():
        return ExtractResult(paras, text, "ok", "")
    return ExtractResult([], "", "partial", "未提取到幻灯片文本")


def _extract_pdf(path: str) -> ExtractResult:
    from pypdf import PdfReader
    reader = PdfReader(path)
    paras = []
    for i, page in enumerate(reader.pages, 1):
        try:
            t = (page.extract_text() or "").strip()
        except Exception:
            t = ""
        if t:
            paras.append("【第%d页】" % i)
            paras.extend([p for p in t.splitlines() if p.strip()])
    text = "\n".join(paras)
    if not text.strip():
        return ExtractResult([], "", "failed", "未提取到文本,疑似扫描版PDF(需OCR,当前版本不支持)")
    if len(text) < 100 and len(reader.pages) >= 2:
        return ExtractResult(paras, text, "partial", "提取文本过少,可能为扫描版PDF")
    return ExtractResult(paras, text, "ok", "")


def extract(path: str) -> ExtractResult:
    """统一入口:按扩展名分发,任何异常都转为 failed 状态,不抛错."""
    ext = os.path.splitext(path)[1].lower()
    try:
        if ext in TEXT_EXTS:
            return _read_text_file(path)
        if ext == ".docx":
            return _extract_docx(path)
        if ext == ".xlsx":
            return _extract_xlsx(path)
        if ext == ".pptx":
            return _extract_pptx(path)
        if ext == ".pdf":
            return _extract_pdf(path)
        return ExtractResult([], "", "failed", "暂不支持的内容格式: %s" % (ext or "(无扩展名)"))
    except Exception as e:
        return ExtractResult([], "", "failed", "解析异常: %s: %s" % (type(e).__name__, e))
