#!/usr/bin/env python3
"""差异层:点开才算.回答"两个版本哪里不一样".

- 文本类(docx/pptx/txt/pdf):按段落 difflib 对齐,返回结构化差异.
- xlsx:单元格级比对,返回 (工作表, 单元格, 旧值, 新值) 清单.
- 明确局限:文本 diff 看不出字体/字号/排版等格式变化.
"""
from __future__ import annotations

import difflib
import html
import os

from . import extract as extmod


def paragraph_diff(paras_a: list, paras_b: list) -> list:
    """返回 [{'tag': 'equal|delete|insert|replace', 'a': [...], 'b': [...]}, ...]"""
    sm = difflib.SequenceMatcher(None, paras_a, paras_b, autojunk=False)
    out = []
    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        out.append({"tag": tag, "a": paras_a[i1:i2], "b": paras_b[j1:j2]})
    return out


def diff_summary(ops: list) -> dict:
    n_eq = sum(len(o["a"]) for o in ops if o["tag"] == "equal")
    n_chg = sum(max(len(o["a"]), len(o["b"])) for o in ops if o["tag"] != "equal")
    total = n_eq + n_chg
    return {"equal": n_eq, "changed": n_chg,
            "similarity": round(n_eq / total, 3) if total else 1.0}


def xlsx_diff(path_a: str, path_b: str, max_rows: int = 2000) -> list:
    """单元格级差异清单."""
    from openpyxl import load_workbook
    wa = load_workbook(path_a, read_only=True, data_only=True)
    wb = load_workbook(path_b, read_only=True, data_only=True)
    diffs = []
    names_a = set(wa.sheetnames)
    names_b = set(wb.sheetnames)
    for s in sorted(names_a - names_b):
        diffs.append({"sheet": s, "cell": "-", "old": "工作表存在", "new": "工作表已删除"})
    for s in sorted(names_b - names_a):
        diffs.append({"sheet": s, "cell": "-", "old": "工作表不存在", "new": "工作表新增"})
    for s in sorted(names_a & names_b):
        sa, sb = wa[s], wb[s]
        for row in sa.iter_rows():
            for c in row:
                if c.row > max_rows:
                    break
                v_old = c.value
                v_new = sb.cell(row=c.row, column=c.column).value
                if v_old != v_new and not (v_old in (None, "") and v_new in (None, "")):
                    diffs.append({"sheet": s, "cell": c.coordinate,
                                  "old": "" if v_old is None else str(v_old),
                                  "new": "" if v_new is None else str(v_new)})
                if len(diffs) > 5000:
                    diffs.append({"sheet": s, "cell": "...", "old": "", "new": "差异过多,仅显示前 5000 条"})
                    wa.close(); wb.close()
                    return diffs
    wa.close(); wb.close()
    return diffs


def diff_files(path_a: str, path_b: str) -> dict:
    ext = os.path.splitext(path_a)[1].lower()
    if ext == ".xlsx" and os.path.splitext(path_b)[1].lower() == ".xlsx":
        d = xlsx_diff(path_a, path_b)
        return {"kind": "xlsx", "diffs": d, "count": len(d)}
    ra, rb = extmod.extract(path_a), extmod.extract(path_b)
    ops = paragraph_diff(ra.paragraphs, rb.paragraphs)
    return {"kind": "text", "ops": ops, "summary": diff_summary(ops),
            "warn": "文本差异;字体/字号/排版等格式变化不在此显示"}


def ops_to_html(ops: list) -> str:
    parts = []
    for o in ops:
        tag = o["tag"]
        if tag == "equal":
            continue
        for p in o["a"]:
            parts.append(f'<div class="del">− {html.escape(p)}</div>')
        for p in o["b"]:
            parts.append(f'<div class="add">＋ {html.escape(p)}</div>')
    return "\n".join(parts) if parts else '<div class="same">两版文本无差异</div>'
